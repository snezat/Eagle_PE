from __future__ import annotations

import copy
import math
import re
import unicodedata
import uuid
from collections import defaultdict
from io import BytesIO
from typing import Any, BinaryIO

from openpyxl import load_workbook


REQUIRED_HEADERS = {
    "Teacher", "Sport", "Last Name", "First Name", "Grade",
    "Original Bench", "New Bench", "Original Squat", "New Squat",
    "Original Power Clean", "New Power Clean",
}

LIFT_COLUMNS = {
    "Bench": ("Original Bench", "New Bench"),
    "Back Squat": ("Original Squat", "New Squat"),
    "Power Clean": ("Original Power Clean", "New Power Clean"),
}


def _name_key(value: str) -> str:
    normalized = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode("ascii")
    return "".join(character for character in normalized.casefold() if character.isalnum())


def _alias_key(value: str) -> str:
    return _name_key(re.sub(r"\([^)]*\)", " ", value))


def _positive_number(value: Any) -> float | None:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return None
    number = float(value)
    if not math.isfinite(number) or number <= 0 or number > 5000:
        return None
    rounded = math.floor(number / 5 + 0.5) * 5
    return int(rounded) if rounded.is_integer() else rounded


def _class_group(sports: list[str]) -> str:
    values = set(sports)
    if "Football" in values:
        return "Football"
    if values.intersection({"Track", "Track & Field", "Cross Country", "Basketball"}):
        return "Nonfootball Group A"
    if values.intersection({"Baseball", "Soccer"}):
        return "Nonfootball Group B"
    return "Nonfootball Group A"


def parse_master_roster(source: bytes | BinaryIO) -> dict[str, Any]:
    stream = BytesIO(source) if isinstance(source, bytes) else source
    try:
        workbook = load_workbook(stream, read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("The uploaded file is not a readable Excel workbook") from exc
    if "Master Roster" not in workbook.sheetnames:
        raise ValueError("The workbook must contain a 'Master Roster' sheet")
    sheet = workbook["Master Roster"]
    header_row = None
    headers: dict[str, int] = {}
    for row_number, row in enumerate(sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 25), values_only=True), start=1):
        candidate = {str(value).strip(): index for index, value in enumerate(row) if value is not None}
        if REQUIRED_HEADERS.issubset(candidate):
            header_row = row_number
            headers = candidate
            break
    if header_row is None:
        missing = ", ".join(sorted(REQUIRED_HEADERS))
        raise ValueError(f"The Master Roster headers were not found. Expected: {missing}")

    grouped: dict[str, dict[str, Any]] = {}
    for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
        first = str(row[headers["First Name"]] or "").strip()
        last = str(row[headers["Last Name"]] or "").strip()
        if not first or not last:
            continue
        name = f"{first} {last}"
        key = _name_key(name)
        athlete = grouped.setdefault(key, {
            "name": name,
            "teachers": set(),
            "grades": set(),
            "sports": set(),
            "newMaxValues": defaultdict(set),
            "originalMaxValues": defaultdict(set),
        })
        teacher = str(row[headers["Teacher"]] or "").strip()
        grade = str(row[headers["Grade"]] or "").strip()
        sport = str(row[headers["Sport"]] or "").strip()
        if teacher:
            athlete["teachers"].add(teacher)
        if grade:
            athlete["grades"].add(grade)
        if sport and sport.casefold() != "unassigned":
            athlete["sports"].add(sport)
        for lift, (original_column, new_column) in LIFT_COLUMNS.items():
            original = _positive_number(row[headers[original_column]])
            current = _positive_number(row[headers[new_column]])
            if original is not None:
                athlete["originalMaxValues"][lift].add(original)
            if current is not None:
                athlete["newMaxValues"][lift].add(current)

    if not grouped:
        raise ValueError("The Master Roster sheet does not contain any athletes")

    athletes = []
    issues = []
    for athlete in grouped.values():
        issue_fields = []
        if len(athlete["teachers"]) > 1:
            issue_fields.append(f"teachers: {', '.join(sorted(athlete['teachers']))}")
        if len(athlete["grades"]) > 1:
            issue_fields.append(f"grades: {', '.join(sorted(athlete['grades']))}")
        for lift, values in athlete["newMaxValues"].items():
            if len(values) > 1:
                issue_fields.append(f"new {lift}: {', '.join(str(value) for value in sorted(values))}")
        if issue_fields:
            issues.append({"name": athlete["name"], "detail": "; ".join(issue_fields)})
        athletes.append({
            "name": athlete["name"],
            "teacher": next(iter(athlete["teachers"]), ""),
            "grade": next(iter(athlete["grades"]), ""),
            "sports": sorted(athlete["sports"]),
            "newMaxes": {lift: next(iter(values)) for lift, values in athlete["newMaxValues"].items() if len(values) == 1},
            "originalMaxes": {lift: next(iter(values)) for lift, values in athlete["originalMaxValues"].items() if len(values) == 1},
            "originalMaxConflicts": {
                lift: sorted(values) for lift, values in athlete["originalMaxValues"].items() if len(values) > 1
            },
        })
    athletes.sort(key=lambda athlete: (_name_key(athlete["name"]), athlete["name"]))
    return {"athletes": athletes, "issues": issues, "sheet": "Master Roster", "headerRow": header_row}


def merge_master_roster(state: dict[str, Any], parsed: dict[str, Any]) -> tuple[dict[str, Any], dict[str, Any]]:
    merged = copy.deepcopy(state)
    current_by_key: dict[str, list[dict[str, Any]]] = defaultdict(list)
    current_by_alias: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for athlete in merged.get("athletes", []):
        current_by_key[_name_key(athlete["name"])].append(athlete)
        current_by_alias[_alias_key(athlete["name"])].append(athlete)
    import_alias_counts: dict[str, int] = defaultdict(int)
    for athlete in parsed["athletes"]:
        import_alias_counts[_alias_key(athlete["name"])] += 1

    matched_ids: set[str] = set()
    added = []
    updated = []
    unchanged = []
    alias_matches = []
    issues = list(parsed.get("issues", []))
    maxes_added = 0
    maxes_updated = 0
    sports_added = 0
    matched = 0

    for imported in parsed["athletes"]:
        exact_candidates = current_by_key.get(_name_key(imported["name"]), [])
        if len(exact_candidates) > 1:
            issues.append({"name": imported["name"], "detail": "More than one stored athlete has this exact name"})
            continue
        current = exact_candidates[0] if exact_candidates else None
        if current is None:
            alias_candidates = [
                athlete for athlete in current_by_alias.get(_alias_key(imported["name"]), [])
                if athlete["id"] not in matched_ids
            ]
            if len(alias_candidates) == 1 and import_alias_counts[_alias_key(imported["name"])] == 1:
                current = alias_candidates[0]
                alias_matches.append({"workbookName": imported["name"], "currentName": current["name"]})
            elif alias_candidates:
                issues.append({"name": imported["name"], "detail": "More than one stored athlete could match this name"})
                continue

        if current is None:
            new_maxes = dict(imported["originalMaxes"])
            new_maxes.update(imported["newMaxes"])
            unresolved = {
                lift: values for lift, values in imported["originalMaxConflicts"].items()
                if lift not in imported["newMaxes"]
            }
            if unresolved:
                details = "; ".join(f"{lift}: {', '.join(str(value) for value in values)}" for lift, values in unresolved.items())
                issues.append({"name": imported["name"], "detail": f"Conflicting original maxes with no newer value ({details})"})
                continue
            current = {
                "id": f"athlete-import-{uuid.uuid4().hex}",
                "name": imported["name"],
                "grade": imported["grade"],
                "teacher": imported["teacher"],
                "classGroup": _class_group(imported["sports"]),
                "sports": list(imported["sports"]),
                "groupBySport": {},
                "subgroup": "",
                "maxes": new_maxes,
                "projectedMaxes": {},
                "overrides": {},
            }
            merged.setdefault("athletes", []).append(current)
            current_by_key[_name_key(current["name"])].append(current)
            matched_ids.add(current["id"])
            added.append(current["name"])
            maxes_added += len(new_maxes)
            sports_added += len(current["sports"])
            continue

        if current["id"] in matched_ids:
            issues.append({"name": imported["name"], "detail": "This workbook row resolves to a student already matched above"})
            continue
        matched += 1
        matched_ids.add(current["id"])
        before = copy.deepcopy(current)
        if imported["grade"]:
            current["grade"] = imported["grade"]
        if imported["teacher"]:
            current["teacher"] = imported["teacher"]
        current_sports = list(dict.fromkeys(current.get("sports", [])))
        combined_sports = list(dict.fromkeys([*current_sports, *imported["sports"]]))
        sports_added += len(set(combined_sports) - set(current_sports))
        current["sports"] = combined_sports
        current["classGroup"] = _class_group(combined_sports)
        current.setdefault("groupBySport", {})
        current["groupBySport"] = {
            sport: group for sport, group in current["groupBySport"].items() if sport in combined_sports
        }
        current.setdefault("maxes", {})
        unresolved = {
            lift: values for lift, values in imported["originalMaxConflicts"].items()
            if lift not in imported["newMaxes"] and lift not in current["maxes"]
        }
        if unresolved:
            details = "; ".join(f"{lift}: {', '.join(str(value) for value in values)}" for lift, values in unresolved.items())
            issues.append({"name": imported["name"], "detail": f"Conflicting original maxes cannot fill an empty stored max ({details})"})
        for lift, value in imported["newMaxes"].items():
            if lift not in current["maxes"]:
                maxes_added += 1
            elif current["maxes"][lift] != value:
                maxes_updated += 1
            current["maxes"][lift] = value
        for lift, value in imported["originalMaxes"].items():
            if lift not in current["maxes"] and lift not in imported["newMaxes"]:
                current["maxes"][lift] = value
                maxes_added += 1
        changes = []
        for field, label in (("grade", "grade"), ("teacher", "teacher"), ("classGroup", "class group"), ("sports", "sports"), ("maxes", "maxes")):
            if before.get(field) != current.get(field):
                changes.append(label)
        if changes:
            updated.append({"name": current["name"], "changes": changes})
        else:
            unchanged.append(current["name"])

    removed = []
    removed_ids: set[str] = set()
    retained_athletes = []
    for athlete in merged.get("athletes", []):
        if athlete["id"] not in matched_ids and athlete.get("classGroup") in {"Nonfootball Group A", "Nonfootball Group B"}:
            removed.append(athlete["name"])
            removed_ids.add(athlete["id"])
        else:
            retained_athletes.append(athlete)
    merged["athletes"] = retained_athletes
    if removed_ids:
        merged["prescriptions"] = [item for item in merged.get("prescriptions", []) if item.get("athleteId") not in removed_ids]
        merged["suggestions"] = [item for item in merged.get("suggestions", []) if item.get("athleteId") not in removed_ids]
        merged["attendance"] = [item for item in merged.get("attendance", []) if item.get("athleteId") not in removed_ids]

    imported_sports = sorted({sport for athlete in parsed["athletes"] for sport in athlete["sports"]})
    merged["sports"] = list(dict.fromkeys([*merged.get("sports", []), *imported_sports]))
    merged.setdefault("sportGroups", {})
    for sport in merged["sports"]:
        merged["sportGroups"].setdefault(sport, [])
    for group in ("Football", "Nonfootball Group A", "Nonfootball Group B"):
        if group not in merged.setdefault("classGroups", []):
            merged["classGroups"].append(group)

    summary = {
        "workbookAthletes": len(parsed["athletes"]),
        "resultingAthletes": len(merged["athletes"]),
        "matched": matched,
        "added": sorted(added),
        "updated": sorted(updated, key=lambda item: item["name"]),
        "unchanged": sorted(unchanged),
        "removed": sorted(removed),
        "aliasMatches": sorted(alias_matches, key=lambda item: item["workbookName"]),
        "maxesAdded": maxes_added,
        "maxesUpdated": maxes_updated,
        "sportsAdded": sports_added,
        "issues": issues,
        "canApply": not issues,
        "revision": state.get("revision", 0),
    }
    return merged, summary
